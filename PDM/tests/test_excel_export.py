"""엑셀 내보내기 테스트 — 문서 종류별 시트 구성, 측정값이 숫자 셀로 들어가는지, 사진 삽입, 표지, 이름 규칙, API 권한"""
import io
import sys
import types
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))
sys.modules.setdefault('eel', types.SimpleNamespace(expose=lambda f: f))

import openpyxl
import pytest
from PIL import Image

from PDM.backend import adapters, excel_export as X, export_api, layout as layout_mod, measurement, runtime
from PDM.backend import api as pdm_api

LAYOUT = layout_mod.effective(None)
TEMPLATES = measurement.load_templates()


def jpeg(color='red', size=(800, 600)):
    buf = io.BytesIO()
    Image.new('RGB', size, color).save(buf, 'JPEG')
    return buf.getvalue()


def load(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def texts(ws):
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value not in (None, '')]


# ── 값 정리 ──────────────────────────────────────────────────────────────────

def test_number_or_text_converts_measurements_to_numbers_and_keeps_text():
    assert X.number_or_text('5') == 5 and isinstance(X.number_or_text('5'), int)
    assert X.number_or_text('-0.05') == -0.05
    assert X.number_or_text('0,5') == 0.5
    assert X.number_or_text('±0.1') == '±0.1'
    assert X.number_or_text('  ') is None and X.number_or_text(None) is None


def test_sheet_name_follows_excel_rules_and_is_unique():
    used = set()
    names = [X.sheet_name('1.크랭크핀/베어링:PORT [A]', used), X.sheet_name('x' * 50, used), X.sheet_name('x' * 50, used), X.sheet_name('', used)]
    assert all(len(n) <= 31 for n in names) and len({n.lower() for n in names}) == 4
    assert not any(ch in names[0] for ch in '\\/*?:[]')


# ── 계측지 ───────────────────────────────────────────────────────────────────

def measurement_doc():
    header = measurement.build_header({'contractNumber': 'SH-1', 'company': '업체', 'shipName': '선박', 'engineModel': 'CUMMINS KTA38'}, '작성자', __import__('datetime').date(2026, 10, 13))
    forms = []
    for form_id, label in (('crankpin', 'PORT'), ('conrod', ''), ('linerpiston', '1호기'), ('camshaft', '')):
        inst = measurement.default_instance(form_id, label)
        inst['phase'] = 'before'
        inst['gaugeNo'] = 'MIC-1'
        if 'tables' in inst:
            for grid in inst['tables'].values():
                grid[0][0], grid[0][1] = '5', '-2'
        else:
            inst['cells'][0][0] = '12'
        forms.append(inst)
    header['forms'] = forms
    return {'docType': 'measurement', 'title': 't', 'header': measurement.normalize_header(header)}


def test_measurement_workbook_has_cover_and_one_sheet_per_form_with_numeric_cells():
    wb = load(X.build_workbook(measurement_doc(), [], None, TEMPLATES, LAYOUT))
    assert wb.sheetnames[0] == '표지' and len(wb.sheetnames) == 5
    cover = texts(wb['표지'])
    assert 'DETAIL REPORT OF ENGINE MAINTENANCE' in cover and any(t.startswith('1. RECORD OF CRANK PIN BEARING(PORT)') for t in cover)
    crank = wb[wb.sheetnames[1]]
    assert 'SH-1' in texts(crank) and 'MIC-1' in texts(crank) and 'BEFORE' in texts(crank)
    numbers = [c.value for row in crank.iter_rows() for c in row if isinstance(c.value, (int, float))]
    assert 5 in numbers and -2 in numbers                       # 측정값은 숫자 셀
    assert any(t.startswith('■ CRANK PIN OUT DIA') and 'Unit' in t for t in texts(crank))


def test_every_builtin_form_exports_without_error_even_when_empty():
    header = measurement.build_header({'contractNumber': 'A'}, 'u', __import__('datetime').date.today())
    header['forms'] = [measurement.default_instance(t['id']) for t in TEMPLATES]
    wb = load(X.build_workbook({'docType': 'measurement', 'title': 't', 'header': measurement.normalize_header(header)}, [], None, TEMPLATES, LAYOUT))
    assert len(wb.sheetnames) == len(TEMPLATES) + 1


def test_conrod_sheet_has_torque_summary_and_position_grid():
    wb = load(X.build_workbook(measurement_doc(), [], None, TEMPLATES, LAYOUT))
    ws = wb[[n for n in wb.sheetnames if '콘로드' in n][0]]
    assert 'D1' in texts(ws) and 'L1' in texts(ws) and 12 in [c.value for row in ws.iter_rows() for c in row]


# ── 서비스레포트 / 부품사진첩 ────────────────────────────────────────────────

def photos_and_images(n, per_page=6):
    photos = [{'id': i + 1, 'slotIndex': slot, 'caption': f'캡션{i + 1}'} for i, slot in enumerate(range(n))]
    return photos, (lambda pid: jpeg('blue', (1600, 1200)))


def sr_doc(per_page=6):
    hd = {'vesselName': '선박', 'imoNo': '123', 'engineType': 'E', 'serialNo': 'S', 'model': 'KTA38', 'date': '2026.10.13', 'orderNo': 'SH-1',
          'reportedBy': '작성자', 'customer': '업체', 'requestedBy': '', 'location': '부산', 'period': '2026.10.01', 'subject': '정비',
          'summary': ['1. 점검', '   - 세부', '2. 교환'], 'remark': '비고1\n비고2', 'pageTitles': {'0': '쿨러 사진'}, 'photosPerPage': per_page}
    return {'docType': 'service_report', 'title': 'sr', 'header': hd}


def test_service_report_has_header_summary_and_photo_sheets_with_images():
    photos, image_of = photos_and_images(8)
    wb = load(X.build_workbook(sr_doc(), photos, image_of, TEMPLATES, LAYOUT))
    assert wb.sheetnames == ['Service report', '사진 1', '사진 2']                    # 6장 + 2장
    body = texts(wb['Service report'])
    assert 'Name of vessel' in body and '선박' in body and '1. 점검' in body and '■ Remark' in body and '비고2' in body
    assert len(wb['사진 1']._images) == 6 and len(wb['사진 2']._images) == 2
    assert '쿨러 사진' in texts(wb['사진 1']) and '캡션1' in texts(wb['사진 1'])


def test_four_per_page_uses_landscape_sheets():
    photos, image_of = photos_and_images(5)
    wb = load(X.build_workbook(sr_doc(4), photos, image_of, TEMPLATES, LAYOUT))
    assert wb['사진 1'].page_setup.orientation == 'landscape' and len(wb['사진 1']._images) == 4 and len(wb['사진 2']._images) == 1


def test_photo_gaps_keep_slot_positions_and_empty_pages_are_skipped():
    photos = [{'id': 1, 'slotIndex': 0, 'caption': 'a'}, {'id': 2, 'slotIndex': 13, 'caption': 'c'}]       # 페이지 0 과 2 (1 은 비어 건너뜀)
    wb = load(X.build_workbook(sr_doc(), photos, lambda _: jpeg(), TEMPLATES, LAYOUT))
    assert wb.sheetnames == ['Service report', '사진 1', '사진 2']
    assert '사진 페이지 3' in texts(wb['사진 2'])


def album_doc(with_cover):
    hd = {'orderNo': 'SALE-1', 'customer': '업체', 'vesselName': '201금성', 'model': 'NO.1 G/E CAT 3412\nNO.2 G/E CUMMINS', 'location': '선진조선',
          'workDate': '2026.07.01', 'withCover': with_cover, 'pageTitles': {'0': 'NO.1 쿨러 작업 사진'}, 'photosPerPage': 6}
    return {'docType': 'photo_album', 'title': 'a', 'header': hd}


def test_album_with_cover_puts_cover_first_with_index_and_place_date():
    photos, image_of = photos_and_images(2)
    wb = load(X.build_workbook(album_doc(True), photos, image_of, TEMPLATES, LAYOUT))
    assert wb.sheetnames == ['표지', '사진 1']
    cover = texts(wb['표지'])
    assert '엔진 작업 사진첩' in cover and '201금성' in cover and '1. NO.1 쿨러 작업 사진 (2매)' in cover
    assert '선진조선' in cover and '2026.07.01' in cover


def test_album_without_cover_has_only_photo_sheets_and_empty_album_still_opens():
    photos, image_of = photos_and_images(1)
    assert load(X.build_workbook(album_doc(False), photos, image_of, TEMPLATES, LAYOUT)).sheetnames == ['사진 1']
    assert len(load(X.build_workbook(album_doc(False), [], None, TEMPLATES, LAYOUT)).sheetnames) == 1


def test_unknown_doc_type_is_rejected():
    with pytest.raises(ValueError):
        X.build_workbook({'docType': 'x', 'title': '', 'header': {}}, [], None, TEMPLATES, LAYOUT)


# ── API ──────────────────────────────────────────────────────────────────────

@pytest.fixture
def db(tmp_path, monkeypatch):
    monkeypatch.setenv('PDM_DB_PATH', str(tmp_path / 'pdm.db'))
    for name in ('_db', '_photos', '_standards'):
        monkeypatch.setattr(runtime, name, None)
    yield runtime.get_db()
    measurement.set_custom_provider(None)


def test_export_api_returns_a_valid_xlsx_named_after_the_document(db):
    doc_id = db.create_document('photo_album', 'S-1', '선박', 'S-1_선박_부품사진첩', album_doc(True)['header'], 'u')
    r = export_api.pdm_export_excel(doc_id)
    import base64
    assert r['success'] and r['filename'] == 'S-1_선박_부품사진첩.xlsx'
    assert load(base64.b64decode(r['data'])).sheetnames[0] == '표지'
    assert export_api.pdm_export_excel(99999)['success'] is False


def test_export_api_is_readable_by_guests_but_still_guarded(db, monkeypatch):
    doc_id = db.create_document('measurement', 'S-1', '선박', 't', measurement_doc()['header'], 'u')
    monkeypatch.setattr(adapters, 'check_access', lambda tier: '로그인이 필요합니다.' if tier != 'guest' else None)
    assert export_api.pdm_export_excel(doc_id)['success'] is True
    monkeypatch.setattr(adapters, 'check_access', lambda tier: '막힘')
    assert export_api.pdm_export_excel(doc_id)['authError'] is True


def test_service_report_header_special_characters_in_title_become_a_safe_filename(db):
    doc_id = db.create_document('service_report', 'S-1', '선박', 'a/b:c*d', sr_doc()['header'], 'u')
    assert export_api.pdm_export_excel(doc_id)['filename'] == 'abcd.xlsx'


def _overlaps(ws):
    ranges = list(ws.merged_cells.ranges)
    return [(a.coord, b.coord) for i, a in enumerate(ranges) for b in ranges[i + 1:]
            if not (a.max_row < b.min_row or b.max_row < a.min_row or a.max_col < b.min_col or b.max_col < a.min_col)]


def test_no_sheet_has_overlapping_merged_ranges_which_excel_reports_as_corrupt():
    header = measurement.build_header({'contractNumber': 'A'}, 'u', __import__('datetime').date.today())
    header['forms'] = [measurement.default_instance(t['id']) for t in TEMPLATES]
    photos, image_of = photos_and_images(7)
    books = [X.build_workbook({'docType': 'measurement', 'title': 't', 'header': measurement.normalize_header(header)}, [], None, TEMPLATES, LAYOUT),
             X.build_workbook(sr_doc(), photos, image_of, TEMPLATES, LAYOUT), X.build_workbook(album_doc(True), photos, image_of, TEMPLATES, LAYOUT)]
    for data in books:
        for ws in load(data).worksheets:
            assert _overlaps(ws) == [], ws.title
