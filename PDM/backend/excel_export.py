"""문서(서비스레포트 / 계측지 / 부품사진첩)를 엑셀(.xlsx)로 내보낸다 — 인쇄 모양을 그대로 옮기는 것이 아니라 엑셀에서 쓰기 좋은 형태.

  · 계측지    : 표지 시트 + 양식 한 장당 시트 하나. 측정값은 숫자 셀로 넣는다 (계산/붙여넣기 편함).
  · 서비스레포트: 헤더/Summary/Remark 시트 + 사진 페이지마다 시트 (사진은 이미지로 삽입)
  · 부품사진첩 : (표지 시트) + 사진 페이지 시트
순수하게 (문서, 사진, 양식 정의, 출력 설정) → xlsx 바이트. DB/Eel 에 의존하지 않는다 (api 쪽에서 읽어 넘긴다).
"""
import io
import re
from typing import Any, Callable, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image, ImageOps

FONT = 'Malgun Gothic'
THIN = Side(style='thin', color='000000')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor='DAEEF3')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)

# 서비스레포트 헤더 항목 (web/sr_fields.js 의 SR_ROWS 와 같다) — [(라벨, 키)] 한 줄 1개(전체 폭) 또는 2개
SR_ROWS = [
    [('Name of vessel', 'vesselName'), ('IMO No.', 'imoNo')], [('Engine Type', 'engineType'), ('S/N', 'serialNo')],
    [('Model', 'model'), ('Date', 'date')], [('Order No.', 'orderNo'), ('Reported by', 'reportedBy')],
    [('Customer', 'customer')], [('Requested by', 'requestedBy')], [('Location', 'location')],
    [('Period', 'period')], [('Subject', 'subject')],
]
PHOTO_DISPLAY = (360, 270)       # 시트에 보이는 사진 크기(px, 4:3)
PHOTO_STORE = (720, 540)         # 파일에 넣는 크기 (표시의 2배 — 확대해도 선명하고 용량은 작게)
_INVALID_SHEET = re.compile(r'[\\/*?:\[\]]')
_INT = re.compile(r'^-?\d+$')
_DECIMAL = re.compile(r'^-?\d+(\.\d+)?$')


def number_or_text(value: Any) -> Any:
    """측정값 문자열 → 숫자(가능하면). 쉼표 소수점 허용. 숫자가 아니면 글자 그대로, 비면 None."""
    text = str(value if value is not None else '').strip()
    if not text:
        return None
    probe = text.replace(',', '.')
    if _INT.match(probe):
        return int(probe)
    return float(probe) if _DECIMAL.match(probe) else text      # nan/inf/1_0 같은 파이썬식 숫자는 글자로 둔다 (엑셀 파일이 깨진다)


def sheet_name(name: str, used: set) -> str:
    """엑셀 시트 이름 규칙(31자, \\ / * ? : [ ] 불가, 중복 불가)에 맞춘다."""
    base = (_INVALID_SHEET.sub('', name).strip().strip("'").strip() or '시트')[:31]
    if base.lower() == 'history':
        base = '시트 History'                                     # 엑셀 예약 이름
    out, n = base, 2
    while out.lower() in used:
        suffix = f' ({n})'
        out, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(out.lower())
    return out


def _cell(ws: Worksheet, ref: str, value: Any = None, *, bold: bool = False, size: float = 10, align: Alignment = CENTER,
          border: bool = False, fill: Optional[PatternFill] = None) -> None:
    c = ws[ref]
    if value is not None:
        c.value = value
        if isinstance(value, str) and value.startswith('='):
            c.data_type = 's'            # '=' 로 시작하는 입력(캡션/요약 등)을 수식으로 실행하지 않고 글자로 둔다
    c.font = Font(name=FONT, size=size, bold=bold)
    c.alignment = align
    if border:
        c.border = BOX
    if fill:
        c.fill = fill


def _merge(ws: Worksheet, rng: str, value: Any = None, **style: Any) -> None:
    ws.merge_cells(rng)
    first = rng.split(':')[0]
    _cell(ws, first, value, **style)
    if style.get('border'):
        for row in ws[rng]:
            for c in row:
                c.border = BOX


def _a4(ws: Worksheet, landscape: bool = False) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4


# ── 사진 페이지 ───────────────────────────────────────────────────────────────

def _photo_bytes(jpeg: bytes, fit: str) -> bytes:
    """칸(4:3)에 맞춘 JPEG. fit == 'cover' 는 가장자리를 잘라 채우고, 'contain' 은 흰 여백으로 전체를 보인다."""
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(jpeg))).convert('RGB')
    img = ImageOps.fit(img, PHOTO_STORE) if fit == 'cover' else ImageOps.pad(img, PHOTO_STORE, color='white')
    buf = io.BytesIO()
    img.save(buf, 'JPEG', quality=85)
    return buf.getvalue()


def _photo_pages(photos: List[Dict[str, Any]], per_page: int) -> List[Dict[str, Any]]:
    """사진이 있는 페이지만 [{page, slots:[사진|None × per_page]}] (칸 위치 유지)."""
    pages: Dict[int, List[Optional[Dict[str, Any]]]] = {}
    for p in photos:
        slots = pages.setdefault(p['slotIndex'] // per_page, [None] * per_page)
        slots[p['slotIndex'] % per_page] = p
    return [{'page': page, 'slots': pages[page]} for page in sorted(pages)]


def add_photo_sheets(wb: Workbook, used: set, header: Dict[str, Any], photos: List[Dict[str, Any]],
                     image_of: Callable[[int], bytes], photo_fit: str) -> List[Dict[str, Any]]:
    """사진 페이지마다 시트 하나 (제목, 2열 사진 + 캡션). 반환: 사진이 있는 페이지 목록 (표지 INDEX 용)."""
    per_page = 4 if header.get('photosPerPage') == 4 else 6
    titles = header.get('pageTitles') or {}
    pages = _photo_pages(photos, per_page)
    for n, pg in enumerate(pages, 1):
        ws = wb.create_sheet(sheet_name(f'사진 {n}', used))
        _a4(ws, landscape=per_page == 4)
        ws.sheet_view.showGridLines = False
        for col, width in zip('ABCDE', (2, 52, 3, 52, 2)):
            ws.column_dimensions[col].width = width
        title = str(titles.get(str(pg['page']), '')).strip()
        ws.row_dimensions[1].height = 26
        _merge(ws, 'B1:D1', title or f'사진 페이지 {pg["page"] + 1}', bold=True, size=13, fill=HEAD_FILL if title else None, border=bool(title))
        for i, photo in enumerate(pg['slots']):
            row = 3 + (i // 2) * 3                                  # 사진 / 캡션 / 여백 세 줄이 한 칸
            col = 'B' if i % 2 == 0 else 'D'
            ws.row_dimensions[row].height = PHOTO_DISPLAY[1] * 0.75    # px → pt
            ws.row_dimensions[row + 1].height = 20
            ws.row_dimensions[row + 2].height = 10
            if not photo:
                continue
            img = XlImage(io.BytesIO(_photo_bytes(image_of(photo['id']), photo_fit)))
            img.width, img.height = PHOTO_DISPLAY
            ws.add_image(img, f'{col}{row}')
            _cell(ws, f'{col}{row + 1}', photo.get('caption') or '', size=10, border=True)
    return pages


# ── 서비스레포트 ──────────────────────────────────────────────────────────────

def _service_report(wb: Workbook, used: set, doc: Dict[str, Any], layout: Dict[str, Any]) -> None:
    hd = doc['header']
    ws = wb.active
    ws.title = sheet_name('Service report', used)
    _a4(ws)
    ws.sheet_view.showGridLines = False
    for col, width in zip('ABCD', (18, 34, 18, 34)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 34
    _merge(ws, 'A1:D1', layout.get('srTitle') or 'Service report', bold=True, size=22)
    row = 3
    for items in SR_ROWS:
        ws.row_dimensions[row].height = 21
        _cell(ws, f'A{row}', items[0][0], bold=True, border=True, fill=HEAD_FILL)
        if len(items) == 1:
            _merge(ws, f'B{row}:D{row}', hd.get(items[0][1], ''), border=True, align=LEFT)
        else:
            _cell(ws, f'B{row}', hd.get(items[0][1], ''), border=True, align=LEFT)
            _cell(ws, f'C{row}', items[1][0], bold=True, border=True, fill=HEAD_FILL)
            _cell(ws, f'D{row}', hd.get(items[1][1], ''), border=True, align=LEFT)
        row += 1
    for heading, lines in (('■ Summary of works', hd.get('summary') or []), ('■ Remark', str(hd.get('remark') or '').split('\n'))):
        row += 1
        _merge(ws, f'A{row}:D{row}', heading, bold=True, size=12, align=LEFT)
        for line in lines:
            row += 1
            indent = len(line) - len(line.lstrip(' '))
            _merge(ws, f'A{row}:D{row}', line.strip(), align=Alignment(horizontal='left', vertical='top', wrap_text=True, indent=min(indent, 15)))
            ws.row_dimensions[row].height = 18 * max(1, len(line.strip()) // 70 + 1)
    row += 2
    _merge(ws, f'A{row}:D{row}', f"{layout.get('company', '')}  {layout.get('address', '')}  TEL {layout.get('tel', '')}  FAX {layout.get('fax', '')}", size=8)


# ── 표지 (부품사진첩 / 계측지) ────────────────────────────────────────────────

def _cover_sheet(wb: Workbook, used: set, name: str, title: str, ship: str, engine: str, index_lines: List[str],
                 extras: List[tuple], layout: Dict[str, Any], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = sheet_name(name, used)
    _a4(ws)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 60
    _merge(ws, 'B2:C2', title, bold=True, size=24)
    ws.row_dimensions[2].height = 50
    _merge(ws, 'B4:C4', ship, bold=True, size=22, fill=HEAD_FILL)
    ws.row_dimensions[4].height = 40
    _merge(ws, 'B5:C5', engine, bold=True, size=16, fill=HEAD_FILL)
    ws.row_dimensions[5].height = 24 * max(1, engine.count('\n') + 1)
    row = 7
    _cell(ws, f'B{row}', 'INDEX', bold=True, size=14, align=LEFT)
    for line in index_lines:
        row += 1
        _merge(ws, f'B{row}:C{row}', line, bold=True, size=12, align=LEFT)
    row += 1
    for label, value in extras:
        row += 1
        _cell(ws, f'B{row}', label, bold=True, size=12, align=LEFT)
        _cell(ws, f'C{row}', value or '', bold=True, size=13, align=LEFT)
    row += 2
    for text in (layout.get('company', ''), layout.get('address', ''), f"Tel {layout.get('tel', '')}   Fax {layout.get('fax', '')}"):
        _merge(ws, f'B{row}:C{row}', text, bold=True, size=11)
        row += 1


def _album(wb: Workbook, used: set, doc, photos, image_of, layout) -> None:
    hd = doc['header']
    placeholder = wb.active
    pages = add_photo_sheets(wb, used, hd, photos, image_of, layout.get('photoFit', 'cover'))
    if hd.get('withCover'):
        lines = []
        titles = hd.get('pageTitles') or {}
        for i, pg in enumerate(pages, 1):
            lines.append(f"{i}. {str(titles.get(str(pg['page']), '')).strip() or '작업 사진'} ({sum(1 for s in pg['slots'] if s)}매)")
        _cover_sheet(wb, used, '표지', '엔진 작업 사진첩', hd.get('vesselName', ''), hd.get('model', ''), lines,
                     [('작업 장소 : ', hd.get('location')), ('작업 일자 : ', hd.get('workDate'))], layout)
        wb.move_sheet(wb.sheetnames[-1], offset=-(len(wb.sheetnames) - 1))     # 표지를 맨 앞으로
    if placeholder.title == 'Sheet' and len(wb.sheetnames) > 1:
        wb.remove(placeholder)


# ── 계측지 ────────────────────────────────────────────────────────────────────

def _measurement_header(ws: Worksheet, form: Dict[str, Any], header: Dict[str, Any], inst: Dict[str, Any], ncols: int) -> int:
    """제목 + 공통 헤더(라벨 2칸 + 값 4칸 짝 둘 × 5줄, 열 12칸 기준). 다음 빈 줄 번호를 돌려준다."""
    ncols = max(ncols, 12)
    ws.row_dimensions[1].height = 28
    _merge(ws, f'A1:{get_column_letter(ncols)}1', form['title'], bold=True, size=15)
    marker = ' '.join(x for x in (header.get('makerType'), inst.get('label')) if x)
    phase = {'before': 'BEFORE', 'after': 'AFTER'}.get(inst.get('phase'), '')
    rows = [
        [('WORK NO', header.get('workNo')), ('INSP. DATE', header.get('inspDate'))],
        [('OWNER', header.get('owner')), ('SHIP NAME', header.get('shipName'))],
        [('MAKER/TYPE', marker), ('INSP. BY', header.get('inspBy'))],
        [('APPR. BY', header.get('apprBy')), ('BEFORE / AFTER', phase)],
        [('계측기기 No.', inst.get('gaugeNo'))],
    ]
    for r, pairs in enumerate(rows, start=3):
        ws.row_dimensions[r].height = 20
        for k, (label, value) in enumerate(pairs):
            c0 = 1 + k * 6
            _merge(ws, f'{get_column_letter(c0)}{r}:{get_column_letter(c0 + 1)}{r}', label, bold=True, border=True, fill=HEAD_FILL)
            _merge(ws, f'{get_column_letter(c0 + 2)}{r}:{get_column_letter(c0 + 5)}{r}', value or '', border=True)
    return 3 + len(rows) + 1


def _journal_sheet(ws: Worksheet, form: Dict[str, Any], header: Dict[str, Any], inst: Dict[str, Any]) -> None:
    max_cols = max(t['cols'] for t in form['tables'])
    ncols = 2 + max_cols
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 8
    for c in range(3, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9
    row = _measurement_header(ws, form, header, inst, max(ncols, 8))
    if form.get('spec'):
        _cell(ws, f'A{row}', 'ITEM', bold=True, border=True, fill=HEAD_FILL)
        _cell(ws, f'B{row}', 'STANDARD', bold=True, border=True, fill=HEAD_FILL, size=8)
        _cell(ws, f'C{row}', 'LIMIT', bold=True, border=True, fill=HEAD_FILL)
        for sp in form['spec']:
            row += 1
            vals = {'standard': '', 'limit': '', **(inst.get('spec') or {}).get(sp['id'], {})}      # 양식이 바뀌어 없는 항목은 빈칸
            _cell(ws, f'A{row}', sp['label'].replace('\n', ' '), border=True, size=9)
            _cell(ws, f'B{row}', number_or_text(vals['standard']), border=True)
            _cell(ws, f'C{row}', number_or_text(vals['limit']), border=True)
        row += 1
    for table in form['tables']:
        row += 1
        unit = ''
        if table.get('unit'):
            unit = f"   Unit : ( {inst['unitBase'].get(table['id'], '')} ) {table['unit']['op']} C/100 mm"
        _merge(ws, f'A{row}:{get_column_letter(ncols)}{row}', f"■ {table['title']}{unit}", bold=True, align=LEFT)
        row += 1
        _merge(ws, f'A{row}:B{row}', f"Position / {table['colLabel'].replace(chr(10), ' ')}", bold=True, border=True, fill=HEAD_FILL, size=8)
        for i in range(table['cols']):
            _cell(ws, f'{get_column_letter(3 + i)}{row}', i + 1, bold=True, border=True, fill=HEAD_FILL)
        data = (inst.get('tables') or {}).get(table['id']) or []
        two = len(table['rows'][0]) == 2
        for ri, labels in enumerate(table['rows']):
            row += 1
            if two:
                _cell(ws, f'A{row}', labels[0] if (table.get('split') or labels[1] == 'TB') else None, bold=True, border=True)
                _cell(ws, f'B{row}', labels[1], border=True)
            else:
                _merge(ws, f'A{row}:B{row}', labels[0], bold=True, border=True)
            for ci, value in enumerate(data[ri] if ri < len(data) else []):
                _cell(ws, f'{get_column_letter(3 + ci)}{row}', number_or_text(value), border=True)
        if two and not table.get('split'):                                # F / A 같은 그룹 라벨은 TB·PS 두 줄에 걸친다
            start = row - len(table['rows']) + 1
            for r0 in range(start, row, 2):
                ws.merge_cells(f'A{r0}:A{r0 + 1}')
        row += 1
    row += 1
    _merge(ws, f'A{row}:{get_column_letter(ncols)}{row}', '■ Remarks :', bold=True, align=LEFT)
    row += 1
    _merge(ws, f'A{row}:{get_column_letter(ncols)}{row + 2}', inst.get('remarks') or '', align=Alignment(horizontal='left', vertical='top', wrap_text=True), border=True)


def _conrod_sheet(ws: Worksheet, form: Dict[str, Any], header: Dict[str, Any], inst: Dict[str, Any]) -> None:
    ncols = 1 + len(form['positions']) * len(form['subColumns'])
    ws.column_dimensions['A'].width = 14
    for c in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 11
    row = _measurement_header(ws, form, header, inst, max(ncols, 8))
    for label, value in (('체결 Torque (kg-m)', inst.get('torque')), ('Original', inst.get('original')), ('STD', inst.get('std')), ('Limit', inst.get('limit'))):
        _merge(ws, f'A{row}:B{row}', label, bold=True, border=True, fill=HEAD_FILL)
        _cell(ws, f'C{row}', number_or_text(value), border=True)
        row += 1
    row += 1
    _merge(ws, f'A{row}:A{row + 1}', "Cyl' No.", bold=True, border=True, fill=HEAD_FILL)
    sub = len(form['subColumns'])
    for i, pos in enumerate(form['positions']):
        _merge(ws, f'{get_column_letter(2 + i * sub)}{row}:{get_column_letter(1 + (i + 1) * sub)}{row}', pos, bold=True, border=True, fill=HEAD_FILL)
        for j, name in enumerate(form['subColumns']):
            _cell(ws, f'{get_column_letter(2 + i * sub + j)}{row + 1}', name, bold=True, border=True, fill=HEAD_FILL)
    row += 1
    for label, cells in zip(inst['rowLabels'], inst['cells']):
        row += 1
        _cell(ws, f'A{row}', label, bold=True, border=True)
        for ci, value in enumerate(cells):
            _cell(ws, f'{get_column_letter(2 + ci)}{row}', number_or_text(value), border=True)


def _measurement(wb: Workbook, used: set, doc: Dict[str, Any], templates: List[Dict[str, Any]], layout: Dict[str, Any]) -> None:
    hd = doc['header']
    by_id = {t['id']: t for t in templates}
    hd = {**hd, 'forms': [f for f in hd.get('forms', []) if f.get('formId') in by_id]}      # 삭제된 사용자 정의 양식은 건너뜀
    def title_of(i: int, inst: Dict[str, Any]) -> str:
        suffix = f"({inst['label']})" if inst.get('label') else ''
        return f"{i}. {by_id[inst['formId']]['indexTitle']}{suffix}"

    lines = [title_of(i, inst) for i, inst in enumerate(hd.get('forms', []), 1)]
    _cover_sheet(wb, used, '표지', 'DETAIL REPORT OF ENGINE MAINTENANCE', hd.get('coverShip', ''), hd.get('coverEngine', ''), lines, [], layout, first=True)
    for i, inst in enumerate(hd.get('forms', []), 1):
        form = by_id[inst['formId']]
        label = f" {inst['label']}" if inst.get('label') else ''
        ws = wb.create_sheet(sheet_name(f"{i}.{form['label']}{label}", used))
        _a4(ws)
        ws.sheet_view.showGridLines = False
        (_conrod_sheet if form['layout'] == 'conrodGrid' else _journal_sheet)(ws, form, hd, inst)
        ws.oddFooter.left.text = form.get('code', '')
        ws.oddFooter.right.text = 'SUNJIN JONGHAP CO., LTD.'


# ── 진입점 ────────────────────────────────────────────────────────────────────

def build_workbook(doc: Dict[str, Any], photos: List[Dict[str, Any]], image_of: Callable[[int], bytes],
                   templates: List[Dict[str, Any]], layout: Dict[str, Any]) -> bytes:
    """doc 종류에 맞는 엑셀 바이트. 알 수 없는 종류는 ValueError."""
    wb = Workbook()
    used: set = set()
    kind = doc['docType']
    if kind == 'service_report':
        _service_report(wb, used, doc, layout)
        add_photo_sheets(wb, used, doc['header'], photos, image_of, layout.get('photoFit', 'cover'))
    elif kind == 'measurement':
        _measurement(wb, used, doc, templates, layout)
    elif kind == 'photo_album':
        _album(wb, used, doc, photos, image_of, layout)
    else:
        raise ValueError(f'엑셀로 내보낼 수 없는 문서 종류: {kind}')
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
